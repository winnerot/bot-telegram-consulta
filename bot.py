import asyncio
import io
import logging
import os
import time
import re
from datetime import datetime
import requests
from bs4 import BeautifulSoup
from telegram import Update
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    MessageHandler,
    ConversationHandler,
    ContextTypes,
    filters,
)
from playwright.async_api import async_playwright
from requests.adapters import HTTPAdapter
from urllib3.util import Retry
import urllib3

# Librería para generar el Word (.docx)
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Configuración de logging y urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO
)
logger = logging.getLogger(__name__)
logging.getLogger("urllib3.connectionpool").setLevel(logging.ERROR)

# Estados de la conversación
ESPERANDO_CAPTCHA = 1

# Constantes SENIAT & Telegram
TELEGRAM_BOT_TOKEN = "8648857231:AAFtDelc_SkKY6-0D1q3_DkH9Hwt7TdgJp0"
SENIAT_USER = "v9684209"
SENIAT_PASS = "v9684209"

# CONFIGURACIÓN DE ADMINISTRADOR Y ACCESO
# Coloca aquí tu ID numérico real de Telegram como Administrador Principal
ADMIN_USER_ID = 2098655970  # <--- REEMPLAZA ESTO CON TU ID REAL

# Archivos de persistencia
SESSION_FILE = "seniat_session.json"
USERS_FILE = "authorized_users.txt"
LOGS_FILE = "consultas_logs.txt"

SESSION_DURATION = 30 * 60  # 30 minutos en segundos
_last_successful_login = 0

# Instancias globales para mantener la sesión persistente de Playwright y la cola de tareas
_playwright_inst = None
_browser_inst = None
_page_inst = None
consulta_lock = asyncio.Lock()  # Cola de tareas para evitar conflictos concurrentes


# ==============================================================================
# GESTIÓN DE ACCESO Y LOGS
# ==============================================================================
def cargar_usuarios_autorizados() -> set:
    autorizados = {ADMIN_USER_ID}
    if os.path.exists(USERS_FILE):
        try:
            with open(USERS_FILE, "r") as f:
                for line in f:
                    uid = line.strip()
                    if uid.isdigit():
                        autorizados.add(int(uid))
        except Exception as e:
            logger.error(f"Error cargando usuarios autorizados: {e}")
    return autorizados


def guardar_usuario_autorizado(user_id: int):
    try:
        autorizados = cargar_usuarios_autorizados()
        autorizados.add(user_id)
        with open(USERS_FILE, "w") as f:
            for uid in autorizados:
                f.write(f"{uid}\n")
    except Exception as e:
        logger.error(f"Error guardando usuario autorizado: {e}")


def remover_usuario_autorizado(user_id: int):
    try:
        autorizados = cargar_usuarios_autorizados()
        if user_id in autorizados:
            autorizados.remove(user_id)
        with open(USERS_FILE, "w") as f:
            for uid in autorizados:
                if uid != ADMIN_USER_ID:
                    f.write(f"{uid}\n")
    except Exception as e:
        logger.error(f"Error removiendo usuario autorizado: {e}")


def registrar_log_consulta(user, cedula: str):
    try:
        ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        username = f"@{user.username}" if user.username else "Sin username"
        nombre_completo = f"{user.first_name or ''} {user.last_name or ''}".strip()
        log_line = f"[{ahora}] | ID: {user.id} | User: {username} | Nombre: {nombre_completo} | Cédula: V{cedula}\n"
        with open(LOGS_FILE, "a", encoding="utf-8") as f:
            f.write(log_line)
    except Exception as e:
        logger.error(f"Error guardando log de consulta: {e}")


# ==============================================================================
# COMANDOS DE ADMINISTRADOR
# ==============================================================================
async def cmd_autorizar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_USER_ID:
        await update.message.reply_text("❌ No tienes permisos de administrador.")
        return

    args = context.args
    target_id = None
    if update.message.reply_to_message:
        target_id = update.message.reply_to_message.from_user.id
    elif args and args[0].isdigit():
        target_id = int(args[0])

    if not target_id:
        await update.message.reply_text("⚠️ Uso incorrecto. Responde a un mensaje o usa: `/autorizar ID`", parse_mode="Markdown")
        return

    guardar_usuario_autorizado(target_id)
    await update.message.reply_text(f"✅ ¡Usuario `{target_id}` autorizado exitosamente!", parse_mode="Markdown")


async def cmd_revocar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_USER_ID:
        await update.message.reply_text("❌ No tienes permisos de administrador.")
        return

    args = context.args
    target_id = None
    if update.message.reply_to_message:
        target_id = update.message.reply_to_message.from_user.id
    elif args and args[0].isdigit():
        target_id = int(args[0])

    if not target_id:
        await update.message.reply_text("⚠️ Uso incorrecto. Responde a un mensaje o usa: `/revocar ID`", parse_mode="Markdown")
        return

    if target_id == ADMIN_USER_ID:
        await update.message.reply_text("❌ No puedes revocarle el acceso al administrador principal.")
        return

    remover_usuario_autorizado(target_id)
    await update.message.reply_text(f"🚫 Acceso revocado para el usuario `{target_id}`.", parse_mode="Markdown")


async def cmd_logs(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_USER_ID:
        await update.message.reply_text("❌ No tienes permisos.")
        return

    if not os.path.exists(LOGS_FILE):
        await update.message.reply_text("ℹ️ Aún no hay registros de consultas guardados.")
        return

    try:
        with open(LOGS_FILE, "rb") as f:
            await update.message.reply_document(document=f, caption="📊 *Historial de Consultas Realizadas*", parse_mode="Markdown")
    except Exception as e:
        await update.message.reply_text(f"❌ Error al enviar logs: {e}")


# ==============================================================================
# GENERADOR DE REPORTE WORD (.DOCX)
# ==============================================================================
def generar_documento_resultado(cedula: str, nombre: str, texto_completo_markdown: str) -> str:
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    COLOR_PRIMARY = RGBColor(31, 78, 121)
    COLOR_SECONDARY = RGBColor(89, 89, 89)
    COLOR_TEXT = RGBColor(51, 51, 51)
    
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titulo.paragraph_format.space_after = Pt(2)
    
    run_sub = p_titulo.add_run("REPORTE DE CONSULTA INTEGRAL\n")
    run_sub.font.size = Pt(9)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_SECONDARY
    
    run_tit = p_titulo.add_run("SENIAT & IVSS - VENEZUELA")
    run_tit.font.size = Pt(16)
    run_tit.font.bold = True
    run_tit.font.color.rgb = COLOR_PRIMARY

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line_run = p_line.add_run("―" * 45)
    p_line_run.font.color.rgb = COLOR_SECONDARY
    p_line.paragraph_format.space_after = Pt(10)

    lineas = texto_completo_markdown.split("\n")
    for linea in lineas:
        linea_limpia = linea.replace("*", "").replace("`", "").strip()
        if not linea_limpia:
            continue
            
        if any(sec in linea_limpia for sec in ["RESULTADO INTEGRAL", "SENIAT", "IVSS", "Contribuyentes", "Empresa"]):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(linea_limpia)
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_PRIMARY
        elif linea_limpia.startswith("•"):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.add_run(linea_limpia.replace("•", "").strip())
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run(linea_limpia)

    filename = f"Resultado_Consulta_V{cedula}.docx"
    doc.save(filename)
    return filename


# ==============================================================================
# PLAYWRIGHT CORE FUNCTIONS
# ==============================================================================
async def get_or_create_page():
    global _playwright_inst, _browser_inst, _page_inst
    if not _playwright_inst:
        _playwright_inst = await async_playwright().start()
    if not _browser_inst or not _browser_inst.is_connected():
        _browser_inst = await _playwright_inst.chromium.launch(headless=True)
    
    if _page_inst and not _page_inst.is_closed():
        return _page_inst

    if os.path.exists(SESSION_FILE):
        context = await _browser_inst.new_context(
            storage_state=SESSION_FILE,
            viewport={"width": 1280, "height": 800}
        )
    else:
        context = await _browser_inst.new_context(viewport={"width": 1280, "height": 800})
        
    _page_inst = await context.new_page()
    return _page_inst


async def guardar_estado_sesion(page):
    global _last_successful_login
    try:
        await page.context.storage_state(path=SESSION_FILE)
        _last_successful_login = time.time()
        print("-> [✓] Sesión y cookies guardadas exitosamente.")
    except Exception as e:
        logger.error(f"No se pudo guardar el estado de la sesión: {e}")


# ==============================================================================
# CONSULTAS PÚBLICAS (SENIAT EMPRESAS & IVSS)
# ==============================================================================
def consultar_seniat_y_nombre(cedula: str):
    url = f"http://contribuyente.seniat.gob.ve/relacionesrif/inicioConsulta.do?personalidad=1&ci={cedula}"
    try:
        response = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=12, verify=False)
        if response.status_code == 200:
            soup = BeautifulSoup(response.text, "html.parser")
            texto_total = soup.get_text(separator=" ", strip=True)

            if any(kw in texto_total.lower() for kw in ["no existe", "no se encuentra", "error", "invalida"]):
                return "No especificado", "❌ *SENIAT (EMPRESA):* La cédula no se encuentra registrada."

            nombre_persona = "No especificado"
            for td in soup.find_all("td"):
                texto_td = td.get_text(strip=True)
                if "NOMBRE:" in texto_td.upper() or "NOMBRE" in texto_td.upper():
                    cleaned = re.sub(r"nombre\s*:\s*", "", texto_td, flags=re.IGNORECASE).strip()
                    if cleaned:
                        nombre_persona = cleaned
                        break

            relaciones = []
            for table in soup.find_all("table"):
                for row in table.find_all("tr"):
                    cols = row.find_all(["td", "th"])
                    cols_text = [c.get_text(strip=True) for c in cols]
                    if len(cols_text) >= 3:
                        primer_col = cols_text[0].upper()
                        if any(primer_col.startswith(pref) for pref in ["V", "J", "E", "G"]) or ("-" in primer_col):
                            relaciones.append({
                                "rif": cols_text[0],
                                "nombre": cols_text[1] if len(cols_text) > 1 else "N/A",
                                "tipo": cols_text[2] if len(cols_text) > 2 else "N/A",
                                "fecha": cols_text[3] if len(cols_text) > 3 else "N/A",
                            })

            resultado_empresas = "🏢 *SENIAT (EMPRESAS / RELACIONES)*\n"
            if relaciones:
                resultado_empresas += "📋 *Contribuyentes Relacionados:*\n"
                for rel in relaciones:
                    resultado_empresas += f"  • *RIF:* `{rel['rif']}`\n"
                    resultado_empresas += f"    *Nombre:* {rel['nombre']}\n"
                    resultado_empresas += f"    *Relación:* {rel['tipo']}\n"
                    resultado_empresas += f"    *Fecha Inicio:* {rel['fecha']}\n"
            else:
                resultado_empresas += "ℹ️ *No se encontraron empresas o relaciones asociadas.*\n"
            
            return nombre_persona, resultado_empresas.strip()
        
        return "No especificado", "❌ *SENIAT (EMPRESA):* Error de conexión."
    except Exception as e:
        logger.error(f"Error SENIAT: {e}")
        return "No especificado", "❌ *SENIAT (EMPRESA):* Error en la petición HTTP."


def consultar_ivss(cedula: str) -> str:
    url_base = "http://www.ivss.gob.ve:28083/CuentaIndividualIntranet/"
    url_accion = "http://www.ivss.gob.ve:28083/CuentaIndividualIntranet/CtaIndividualCTRL"
    
    headers = {
        "User-Agent": "Mozilla/5.0",
        "Referer": url_base,
        "Origin": "http://www.ivss.gob.ve:28083/CuentaIndividualIntranet/",
        "Content-Type": "application/x-www-form-urlencoded",
        "Connection": "close",
    }
    
    try:
        session = requests.Session()
        retries = Retry(total=3, backoff_factor=1, status_forcelist=[500, 502, 503, 504], allowed_methods=["GET", "POST"])
        adapter = HTTPAdapter(max_retries=retries)
        session.mount("http://", adapter)
        session.mount("https://", adapter)

        session.get(url_base, headers=headers, timeout=15, verify=False)
        payload = {"Accion": "", "Accion1": "", "nacionalidad_aseg": "V", "cedula_aseg": cedula, "consultar": "Buscar"}
        response = session.post(url_accion, data=payload, headers=headers, timeout=20, verify=False)

        if response.status_code == 200:
            soup = BeautifulSoup(response.text, "html.parser")
            texto_total = soup.get_text(separator=" ", strip=True)
            if any(kw in texto_total.lower() for kw in ["no existe", "no registrado", "error"]):
                return f"🛡️ *IVSS SEGURO:* La cédula `{cedula}` no se encuentra registrada."

            empresas, fechas_ingreso, estatus_list = [], [], []
            for tr in soup.find_all("tr"):
                tds = tr.find_all(["td", "th"])
                for i, td in enumerate(tds):
                    txt = td.get_text(strip=True).upper()
                    if "EMPRESA" in txt or "RAZON SOCIAL" in txt:
                        if i + 1 < len(tds):
                            val = tds[i + 1].get_text(strip=True)
                            if val and val != ":" and val not in empresas:
                                empresas.append(val)
                    if "FECHA" in txt and "INGRESO" in txt:
                        if i + 1 < len(tds):
                            val = tds[i + 1].get_text(strip=True)
                            if val and val != ":" and val not in fechas_ingreso:
                                fechas_ingreso.append(val)
                    if "ESTATUS" in txt or "ESTADO" in txt:
                        if i + 1 < len(tds):
                            val = tds[i + 1].get_text(strip=True)
                            if val and val != ":" and val not in estatus_list:
                                estatus_list.append(val)

            resultado_msj = "🛡️ *IVSS SEGURO*\n"
            if empresas or fechas_ingreso or estatus_list:
                if empresas:
                    resultado_msj += "*Empresa(s):*\n" + "\n".join([f"• {e}" for e in empresas]) + "\n"
                if fechas_ingreso:
                    resultado_msj += "*Fecha(s) de Ingreso:*\n" + "\n".join([f"• {f}" for f in fechas_ingreso]) + "\n"
                if estatus_list:
                    resultado_msj += "*Estatus:*\n" + "\n".join([f"• {es}" for es in estatus_list]) + "\n"
                return resultado_msj.strip()
            return "🛡️ *IVSS SEGURO:* NO ASEGURADO"
        
        return "❌ *IVSS SEGURO:* El servicio tardó demasiado en responder."
    except Exception as e:
        logger.warning(f"Aviso IVSS: {e}")
        return "⚠️ *IVSS SEGURO:* El servidor tardó en responder."


# ==============================================================================
# FUNCIÓN CENTRAL DE BÚSQUEDA Y EXTRACCIÓN + ENVÍO DE DOCX
# ==============================================================================
async def ejecutar_consulta_rif(page, cedula_target, update: Update, msg_espera):
    await page.goto("http://contribuyente.seniat.gob.ve/rifconsulta/login.do", timeout=30000)
    await page.wait_for_load_state("networkidle")
    
    await page.wait_for_selector("input[name='cedula']", timeout=15000)
    await page.fill("input[name='cedula']", f"V{cedula_target}")
    await page.click("input[name='buscar']")
    await asyncio.sleep(4)

    enlace_expediente = page.locator(f"a[href*='buscarexpediente.do']:has-text('{cedula_target}')")
    if await enlace_expediente.count() == 0:
        enlace_expediente = page.locator("a[href*='buscarexpediente.do']").first
    
    await enlace_expediente.click()
    await asyncio.sleep(3)

    enlace_direcciones = page.locator("a.menuitem[href*='datosdirecciones.do']").first
    await enlace_direcciones.click()
    await asyncio.sleep(3)

    enlace_domicilio = page.locator("a[href*='datosdirecciones.do?consultar=S']").first
    await enlace_domicilio.click()
    await asyncio.sleep(3)

    dir_info = await page.evaluate(r"""() => {
        const tables = document.querySelectorAll('table');
        let t = Array.from(tables).find(tbl => tbl.innerText.includes('Tipo de Dirección') && tbl.innerText.includes('Estado'));
        if (!t) return null;
        const trs = t.querySelectorAll('tr');
        let resultados = [];
        for (let i = 0; i < trs.length - 1; i++) {
            let currentTr = trs[i];
            let nextTr = trs[i+1];
            let subs = Array.from(currentTr.querySelectorAll('.tablaSubTitulo')).map(e => e.innerText.replace(/\u00a0/g, ' ').trim());
            let smalls = Array.from(nextTr.querySelectorAll('.letrasSmall')).map(e => e.innerText.replace(/\u00a0/g, ' ').trim());
            if (subs.length > 0 && smalls.length > 0) {
                for (let j = 0; j < subs.length; j++) {
                    let header = subs[j] ? subs[j] : '';
                    let val = smalls[j] || 'NO INDICA';
                    if (header) {
                        resultados.push({ etiqueta: header, valor: val.toUpperCase() });
                    }
                }
            }
        }
        return resultados;
    }""")

    nombre_persona, res_seniat_empresa = await asyncio.to_thread(consultar_seniat_y_nombre, cedula_target)
    res_ivss = await asyncio.to_thread(consultar_ivss, cedula_target)
    
    mensaje_final = f"📋 *RESULTADO INTEGRAL DE CONSULTA*\n\n"
    mensaje_final += f"👤 *Cédula:* `V{cedula_target}`\n"
    mensaje_final += f"👤 *Nombre:* {nombre_persona}\n\n"
    
    if dir_info:
        mensaje_final += "🏢 *SENIAT (DOMICILIO FISCAL)*\n"
        for item in dir_info:
            etiqueta = item.get("etiqueta", "").title()
            valor = item.get("valor", "NO INDICA")
            if etiqueta.strip() and valor.strip():
                mensaje_final += f"• *{etiqueta}:* {valor}\n"
        mensaje_final += "\n"
    else:
        mensaje_final += "🏢 *SENIAT (DOMICILIO FISCAL):* No se pudo extraer la información detallada.\n\n"

    mensaje_final += res_seniat_empresa + "\n\n"
    mensaje_final += res_ivss

    # Borrar mensaje de espera
    try:
        await update.get_bot().delete_message(chat_id=update.effective_chat.id, message_id=msg_espera.message_id)
    except Exception:
        pass

    # 1. Enviar el mensaje estructurado de texto por Telegram
    await update.message.reply_text(mensaje_final, parse_mode="Markdown")

    # 2. Generar y enviar el documento Word (.docx) automáticamente
    try:
        archivo_docx = await asyncio.to_thread(generar_documento_resultado, cedula_target, nombre_persona, mensaje_final)
        with open(archivo_docx, "rb") as doc_file:
            await update.message.reply_document(
                document=doc_file,
                caption=f"📄 *Reporte en Word generado para la cédula V{cedula_target}*",
                parse_mode="Markdown"
            )
        # Limpiar archivo local
        if os.path.exists(archivo_docx):
            os.remove(archivo_docx)
    except Exception as e:
        logger.error(f"Error generando o enviando el archivo .docx: {e}")


# ==============================================================================
# MANEJADOR DE MENSAJES Y ACCESO
# ==============================================================================
async def manejar_cedula_o_texto(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    usuarios_permitidos = cargar_usuarios_autorizados()
    
    if user.id not in usuarios_permitidos:
        await update.message.reply_text(
            "⛔ *Acceso Denegado*\n\nNo estás autorizado para utilizar este bot.\n"
            f"Tu ID de Telegram es: `{user.id}`\nComunícate con el administrador para solicitar acceso.",
            parse_mode="Markdown"
        )
        return ConversationHandler.END

    raw_text = update.message.text.strip()
    cedula_digits = re.sub(r"\D", "", raw_text)

    if not cedula_digits or len(cedula_digits) < 6 or len(cedula_digits) > 9:
        await update.message.reply_text("👋 Envíame un número de cédula válido para consultar.")
        return ConversationHandler.END

    # Registrar Log de la consulta
    registrar_log_consulta(user, cedula_digits)

    msg_espera = await update.message.reply_text("🔄 Procesando consulta... (Esperando turno en cola).")

    async with consulta_lock:
        try:
            page = await get_or_create_page()
            tiempo_transcurrido = time.time() - _last_successful_login
            sesion_vigente = tiempo_transcurrido < SESSION_DURATION

            necesita_login = True
            if sesion_vigente:
                try:
                    await page.goto("http://contribuyente.seniat.gob.ve/rifconsulta/login.do", timeout=15000)
                    await page.wait_for_load_state("networkidle", timeout=8000)
                    if "iseniatlogin" not in page.url and await page.locator("input[name='cedula']").count() > 0:
                        necesita_login = False
                except Exception:
                    necesita_login = True

            if necesita_login:
                await page.goto("http://contribuyente.seniat.gob.ve/iseniatlogin/contribuyente.do", timeout=60000)
                await page.wait_for_load_state("networkidle")

                captcha_img = page.locator("img[src*='kapatch']")
                await captcha_img.wait_for(timeout=10000)
                captcha_bytes = await captcha_img.screenshot()

                context.user_data["pending_cedula"] = cedula_digits

                try:
                    await update.get_bot().delete_message(chat_id=update.effective_chat.id, message_id=msg_espera.message_id)
                except Exception:
                    pass

                file_stream = io.BytesIO(captcha_bytes)
                file_stream.name = "captcha_seniat.jpg"

                sent_msg = await update.message.reply_photo(
                    photo=file_stream,
                    caption=f"🧩 *Sesión expirada.* \n\nPor favor, **responde a este mensaje** con los caracteres del captcha para continuar con la cédula V{cedula_digits}:",
                    parse_mode="Markdown"
                )
                context.user_data["captcha_msg_id"] = sent_msg.message_id
                return ESPERANDO_CAPTCHA
            else:
                await ejecutar_consulta_rif(page, cedula_digits, update, msg_espera)
                return ConversationHandler.END

        except Exception as e:
            logger.error(f"Error: {e}")
            try:
                await update.get_bot().delete_message(chat_id=update.effective_chat.id, message_id=msg_espera.message_id)
            except Exception:
                pass
            await update.message.reply_text("❌ Ocurrió un error vuelve a intentarlo.")
            return ConversationHandler.END


async def recibir_captcha_y_consultar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    codigo_usuario = update.message.text.strip()
    cedula_target = context.user_data.get("pending_cedula")
    captcha_msg_id = context.user_data.get("captcha_msg_id")

    if captcha_msg_id:
        try:
            await update.get_bot().delete_message(chat_id=update.effective_chat.id, message_id=captcha_msg_id)
        except Exception:
            pass

    try:
        await update.message.delete()
    except Exception:
        pass

    if not cedula_target:
        await update.message.reply_text("⚠️ Sesión expirada. Envía la cédula nuevamente.")
        return ConversationHandler.END

    msg_espera = await update.message.reply_text("🔄 Procesando consulta...")

    async with consulta_lock:
        try:
            page = await get_or_create_page()
            await page.fill("input[name='usuario']", SENIAT_USER)
            await page.fill("input[name='clave']", SENIAT_PASS)
            await page.fill("input[name='kaptcha']", codigo_usuario)

            await page.evaluate("validarIngreso()")
            await asyncio.sleep(4) 

            error_locator = page.locator("#mensajeError")
            error_text = await error_locator.inner_text() if await error_locator.count() > 0 else ""
            
            if error_text.strip() and "incorrecto" in error_text.lower():
                try:
                    await update.get_bot().delete_message(chat_id=update.effective_chat.id, message_id=msg_espera.message_id)
                except Exception:
                    pass
                await update.message.reply_text(f"❌ El SENIAT rechazó el acceso: _{error_text.strip()}_. Envía la cédula de nuevo.", parse_mode="Markdown")
                context.user_data.clear()
                return ConversationHandler.END

            await guardar_estado_sesion(page)
            await ejecutar_consulta_rif(page, cedula_target, update, msg_espera)

        except Exception as e:
            logger.error(f"Error tras captcha: {e}")
            try:
                await update.get_bot().delete_message(chat_id=update.effective_chat.id, message_id=msg_espera.message_id)
            except Exception:
                pass
            await update.message.reply_text("❌ Ocurrió un error vuelve a intentarlo.")
        finally:
            context.user_data.clear()

    return ConversationHandler.END


async def cancelar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("❌ Operación cancelada.")
    return ConversationHandler.END


def main():
    app = ApplicationBuilder().token(TELEGRAM_BOT_TOKEN).build()

    # Comandos de Administrador
    app.add_handler(CommandHandler("autorizar", cmd_autorizar))
    app.add_handler(CommandHandler("revocar", cmd_revocar))
    app.add_handler(CommandHandler("logs", cmd_logs))

    # Conversación de Consultas
    conv_handler = ConversationHandler(
        entry_points=[MessageHandler(filters.TEXT & (~filters.COMMAND), manejar_cedula_o_texto)],
        states={
            ESPERANDO_CAPTCHA: [MessageHandler(filters.TEXT & (~filters.COMMAND), recibir_captcha_y_consultar)]
        },
        fallbacks=[CommandHandler("cancelar", cancelar)],
    )

    app.add_handler(conv_handler)

    print("🤖 Bot con control de acceso, logs y reportes Word activo...")
    app.run_polling(bootstrap_retries=-1)


if __name__ == "__main__":
    main()
